"""Prove the inspector can fail: build broken copies of the decks in scratch/ and expect specific failures.

Cases:
  1. one charted value changed in the chart cache (HUD O&M 97 -> 96), workbook left alone
  2. a picture pasted onto the faithful slide
  3. value labels switched on for every bar of the redesign
  4. an en dash pasted into the redesign's subtitle
  5. file properties that claim the deck came from Microsoft Office PowerPoint
Each broken folder must make src/inspect_decks.py exit non-zero and name the expected check as FAIL.
The untouched decks must pass.

Usage: python3 src/test_inspector.py
"""

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
GOOD = ROOT / "slides" / "figure1"
WORK = ROOT / "scratch" / "inspector-tests"
PY = sys.executable


def run(folder: Path) -> tuple[int, str]:
    p = subprocess.run([PY, str(ROOT / "src" / "inspect_decks.py"), str(folder)], capture_output=True, text=True)
    return p.returncode, p.stdout


def copy_case(name: str) -> Path:
    d = WORK / name
    d.mkdir(parents=True, exist_ok=True)
    for f in ("faithful.pptx", "redesign.pptx"):
        shutil.copy(GOOD / f, d / f)
    return d


def rewrite_part(pptx: Path, part: str, edit) -> None:
    src = zipfile.ZipFile(pptx)
    items = {n: src.read(n) for n in src.namelist()}
    src.close()
    items[part] = edit(items[part].decode("utf-8")).encode("utf-8")
    with zipfile.ZipFile(pptx, "w", zipfile.ZIP_DEFLATED) as out:
        for n, data in items.items():
            out.writestr(n, data)


def main() -> int:
    failures = []
    code, out = run(GOOD)
    if code != 0:
        failures.append("the real decks do not pass")

    d = copy_case("changed-value")
    # HUD is figure row 18, point index 17, in the O&M series cache.
    def change(xml: str) -> str:
        new, n = re.subn(r'(<c:pt idx="17"><c:v>)97(</c:v>)', r"\g<1>96\g<2>", xml, count=1)
        assert n == 1, "HUD value not found"
        return new
    rewrite_part(d / "faithful.pptx", "ppt/charts/chart1.xml", change)
    code, out = run(d)
    if code == 0 or "FAIL  faithful.pptx  series 'O&M' values match the CSV" not in out:
        failures.append("a changed chart value was not caught")

    d = copy_case("pasted-picture")
    prs = Presentation(str(d / "faithful.pptx"))
    prs.slides[0].shapes.add_picture(str(GOOD / "source.png"), Inches(0.2), Inches(0.2), Inches(2))
    prs.save(str(d / "faithful.pptx"))
    code, out = run(d)
    if code == 0 or "FAIL  faithful.pptx  no pictures on the slide" not in out:
        failures.append("a pasted picture was not caught")

    d = copy_case("all-labels")
    rewrite_part(d / "redesign.pptx", "ppt/charts/chart1.xml",
                 lambda xml: xml.replace('<c:showLegendKey val="0"/><c:showVal val="0"/>', '<c:showLegendKey val="0"/><c:showVal val="1"/>', 1))
    code, out = run(d)
    if code == 0 or "FAIL  redesign.pptx  value labels only on the five extremes" not in out:
        failures.append("labels on every bar were not caught")

    d = copy_case("dash-in-text")
    dash = chr(0x2013)
    rewrite_part(d / "redesign.pptx", "ppt/slides/slide1.xml",
                 lambda xml: xml.replace("operations and maintenance (O&amp;M)", f"operations {dash} maintenance (O&amp;M)", 1))
    code, out = run(d)
    if code == 0 or "FAIL  redesign.pptx  no dash characters or banned words" not in out:
        failures.append("a dash in slide text was not caught")

    d = copy_case("claims-powerpoint")
    rewrite_part(d / "faithful.pptx", "docProps/app.xml",
                 lambda xml: re.sub(r"<Application>[^<]*</Application>", "<Application>Microsoft Office PowerPoint</Application>", xml))
    code, out = run(d)
    if code == 0 or "FAIL  faithful.pptx  file properties say how the deck was made" not in out:
        failures.append("properties claiming PowerPoint were not caught")

    for f in failures:
        print("TEST FAILED:", f)
    print("inspector tests:", "all 6 behave as expected" if not failures else f"{len(failures)} problem(s)")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
